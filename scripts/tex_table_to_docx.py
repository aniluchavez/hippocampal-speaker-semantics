#!/usr/bin/env python3
"""Convert one of this project's results/*.tex tables into a publication-style
.docx: bold "Table N. <caption>" title, a three-line (APA-style) table with
no vertical gridlines and no rules between data rows, bold centered header,
numeric columns centered / label columns left-aligned, grouped rows (where
the LaTeX has a blank first cell / \\addlinespace to indicate a continuation
of the previous group, e.g. multiple tests per model) vertically merged in
the first column, and the \\flushleft\\footnotesize footnote rendered below
the table in small type.

This is a lightweight regex-based delatexer scoped to the macros actually
used in this project's generated tables (\\mathrm, \\text, \\pm, \\times,
\\cdot, \\sqrt, \\%, \\ref, \\min, ~, --/---, `` '' quotes, sub/superscripts,
\\resizebox, \\addlinespace) -- not a general LaTeX parser.

Usage:
    python3 scripts/tex_table_to_docx.py results/some_table.tex [output.docx] [--number N]
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def extract_braced(text: str, start: int) -> tuple[str, int]:
    """Given text[start] == '{', return (inner_content, index_after_closing_brace)."""
    assert text[start] == "{"
    depth = 0
    for i in range(start, len(text)):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:i], i + 1
    raise ValueError("unbalanced braces")


def extract_command_arg(text: str, command: str) -> str | None:
    """Find \\command{...} and return its (brace-balanced) contents, or None."""
    m = re.search(re.escape(command) + r"\{", text)
    if not m:
        return None
    inner, _ = extract_braced(text, m.end() - 1)
    return inner


def strip_wrapping_commands(text: str) -> str:
    """Repeatedly strip \\mathrm{X}, \\text{X}, \\emph{X}, \\textbf{X} -> X."""
    pattern = re.compile(r"\\(mathrm|text|emph|textbf|textit|mathbf)\{")
    while True:
        m = pattern.search(text)
        if not m:
            break
        inner, end = extract_braced(text, m.end() - 1)
        text = text[:m.start()] + inner + text[end:]
    return text


SUPERSCRIPT_DIGITS = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")
SUBSCRIPT_CHARS = str.maketrans("0123456789+-", "₀₁₂₃₄₅₆₇₈₉₊₋")


def delatex(text: str) -> str:
    """Best-effort LaTeX -> plain text for the macro subset used in these tables."""
    text = strip_wrapping_commands(text)

    # \ref{tab:xyz} -> [xyz] (docx has no cross-reference numbering to resolve to)
    text = re.sub(r"\\ref\{(?:tab:)?([^}]*)\}", r"[\1]", text)
    text = re.sub(r"\\label\{[^}]*\}", "", text)

    def _sup(m):
        exp = m.group(1) or m.group(2)
        return exp.translate(SUPERSCRIPT_DIGITS)
    text = re.sub(r"\^\{([0-9]+)\}|\^([0-9])", _sup, text)

    def _sub(m):
        sub = m.group(1) or m.group(2)
        return sub.translate(SUBSCRIPT_CHARS)
    text = re.sub(r"_\{([0-9+\-]+)\}|_([0-9+\-])", _sub, text)

    # superscript symbols (non-digit) -- currently just \dagger/\ddagger footnote markers
    text = re.sub(r"\^\{?\\ddagger\}?", "‡", text)
    text = re.sub(r"\^\{?\\dagger\}?", "†", text)

    replacements = {
        r"\pm": "±", r"\times": "×", r"\cdot": "·", r"\%": "%", r"\$": "$",
        r"\ll": "«", r"\gg": "»", r"``": "“", r"''": "”",
        r"\-\-\-": "—", r"\-\-": "–", r"~": " ",
        r"\min": "min", r"\max": "max", r"\ddagger": "‡", r"\dagger": "†",
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    text = text.replace("---", "—").replace("--", "–")

    text = re.sub(r"\\sqrt\{([^{}]*)\}", r"sqrt(\1)", text)
    text = text.replace("$", "")
    text = text.replace(r"\_", "_")
    text = re.sub(r"\\ ", " ", text)
    # any remaining unhandled \command (no args) -> drop the backslash
    text = re.sub(r"\\([a-zA-Z]+)", r"\1", text)

    text = re.sub(r"[ \t]+", " ", text).strip()
    return text


def parse_table(tex: str) -> dict:
    caption = extract_command_arg(tex, r"\caption")
    if caption is None:
        raise ValueError("no \\caption{...} found")

    tab_m = re.search(r"\\begin\{tabular\}\{[^}]*\}", tex)
    if not tab_m:
        raise ValueError("no \\begin{tabular}{...} found")
    body_start = tab_m.end()
    body_end = tex.index(r"\end{tabular}", body_start)
    body = tex[body_start:body_end]

    body = re.sub(r"\\(toprule|midrule|bottomrule)", "", body)

    # \addlinespace marks a group boundary -- split on it first so its bare
    # text (no & or \\ of its own) never leaks into a cell.
    segments = re.split(r"\\addlinespace", body)
    table_rows = []
    new_group_flags = []
    for seg_idx, seg in enumerate(segments):
        seg_rows = [r.strip() for r in seg.split(r"\\") if r.strip()]
        for i, raw_row in enumerate(seg_rows):
            cells = [delatex(c) for c in raw_row.split("&")]
            table_rows.append(cells)
            new_group_flags.append(seg_idx > 0 and i == 0)

    footnote = None
    fn_m = re.search(r"\\begin\{flushleft\}(.*?)\\end\{flushleft\}", tex, re.S)
    if fn_m:
        fn_text = fn_m.group(1).replace(r"\footnotesize", "")
        footnote = delatex(fn_text)

    return {
        "caption": delatex(caption),
        "rows": table_rows,
        "new_group_flags": new_group_flags,
        "footnote": footnote,
    }


def _set_cell_borders(cell, top=None, bottom=None):
    """Set only top/bottom borders on a cell (no left/right/inside lines)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        spec = top if side == "top" else bottom if side == "bottom" else None
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, color = spec
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


NUMERIC_RE = re.compile(r"^[\s0-9±×·%\.\-–—+\[\],()a-zA-Z<≈=e]*$")


def _looks_numeric(values: list[str]) -> bool:
    non_empty = [v for v in values if v]
    if not non_empty:
        return False
    return all(NUMERIC_RE.match(v) and any(ch.isdigit() for ch in v) for v in non_empty)


def build_docx(parsed: dict, out_path: Path, table_number: int | None) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    prefix = f"Table {table_number}. " if table_number is not None else ""
    run = title.add_run(prefix + parsed["caption"])
    run.bold = True
    run.font.size = Pt(11)

    rows = parsed["rows"]
    new_group_flags = parsed["new_group_flags"]
    header, data_rows = rows[0], rows[1:]
    data_flags = new_group_flags[1:]
    ncols = len(header)

    # decide per-column alignment: numeric-looking columns centered, else left
    col_is_numeric = [
        _looks_numeric([r[c] for r in data_rows if c < len(r)])
        for c in range(ncols)
    ]
    col_is_numeric[0] = False  # first column is always the row label

    t = doc.add_table(rows=1, cols=ncols)
    t.autofit = True

    THIN = (4, "000000")
    THICK = (12, "000000")

    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        _set_cell_borders(cell, top=THICK, bottom=THIN)

    n_data = len(data_rows)
    for ridx, row in enumerate(data_rows):
        cells = t.add_row().cells
        is_last = ridx == n_data - 1
        for i in range(ncols):
            val = row[i] if i < len(row) else ""
            p = cells[i].paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_is_numeric[i] else WD_ALIGN_PARAGRAPH.LEFT
            )
            if data_flags[ridx] and ridx > 0:
                p.paragraph_format.space_before = Pt(6)
            r = p.add_run(val)
            r.font.size = Pt(10)
            if i == 0 and val:
                r.bold = True
            _set_cell_borders(cells[i], bottom=THICK if is_last else None)

    # merge continuation rows (blank leading cells) into the group row above,
    # for column 0 (row label) and any other leading columns that are also
    # blank on the continuation row (e.g. a per-group summary column)
    group_start_row = [1] * ncols  # row 0 is header
    for ridx, row in enumerate(data_rows):
        table_row_idx = ridx + 1
        if row[0] != "" or table_row_idx == 1:
            for c in range(ncols):
                group_start_row[c] = table_row_idx
            continue
        for c in range(ncols):
            val = row[c] if c < len(row) else ""
            if val != "":
                group_start_row[c] = table_row_idx
                continue
            top_cell = t.cell(group_start_row[c], c)
            merged = top_cell.merge(t.cell(table_row_idx, c))
            merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # drop the now-redundant empty paragraph(s) left by the blank
            # continuation-row cell, keeping only the real label paragraph
            for p in merged.paragraphs[1:]:
                if not p.text.strip():
                    p._element.getparent().remove(p._element)

    if parsed["footnote"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run(parsed["footnote"])
        run.font.size = Pt(8.5)

    doc.save(str(out_path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tex_path")
    ap.add_argument("out_path", nargs="?", default=None)
    ap.add_argument("--number", type=int, default=None, help='Prefix title with "Table N."')
    args = ap.parse_args()

    in_path = Path(args.tex_path)
    out_path = Path(args.out_path) if args.out_path else in_path.with_suffix(".docx")

    tex = in_path.read_text()
    parsed = parse_table(tex)
    build_docx(parsed, out_path, args.number)
    print(f"Saved -> {out_path}")


if __name__ == "__main__":
    main()
